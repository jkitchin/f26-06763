# /// script
# requires-python = ">=3.11"
# dependencies = ["numpy", "matplotlib", "pillow", "python-pptx"]
# ///
"""Figures for Assignment 5.

    uv run --no-project course/assignments/figures/make_a05_figures.py

    a05-reactor.png  Victor Alves's bioreactor drawing (cmu-seminar slide 35, the
                     same picture as Lecture 13's fedbatch-reactor.png), relabeled
                     for continuous operation: feed in, outflow out, constant volume.
                     Needs attic/cmu-seminar.pptx, which is not in the repository;
                     without it the committed PNG is kept.
    a05-data.png     The five experiments of a05-bioreactor.csv: experiments 1 to 4
                     for 100 s (training) and experiment 5 for 300 s (test), with the
                     first 40 s of experiment 5, which Problems 1 and 2 train on, filled.
    a05-physics.png  The two of Lecture 13's ways to add the physics to a network that this
                     homework uses, the PINN and the neural ODE, redrawn from its slide "four
                     ways to add the physics" with its own figures.
"""
import csv
import io
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent.parent
DATA = HERE.parent / "a05-bioreactor.csv"

INK = "#1a1a1a"
MUTED = "#5c5c5c"
BLUE = "#1f5c99"
ORANGE = "#c2410c"
STYLE = {
    "font.size": 14, "axes.labelsize": 14, "axes.titlesize": 15,
    "xtick.labelsize": 13, "ytick.labelsize": 13, "legend.fontsize": 12,
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.edgecolor": MUTED, "text.color": INK, "axes.labelcolor": INK,
    "xtick.color": MUTED, "ytick.color": MUTED, "legend.frameon": False,
    "savefig.dpi": 150, "savefig.bbox": "tight",
}
X_SS, S_SS = 3.46198, 45.6033      # the steady state, kg/m3 and mol/m3


def save(fig, name):
    fig.savefig(HERE / name)
    plt.close(fig)
    print(f"wrote {name}")


def reactor():
    deck = REPO / "attic" / "cmu-seminar.pptx"
    if not deck.exists():
        print(f"skipped a05-reactor.png: {deck} is not here")
        return
    from PIL import Image, ImageDraw
    from pptx import Presentation

    shape = next(s for s in Presentation(str(deck)).slides[34].shapes if s.name == "Picture 17")
    img = Image.open(io.BytesIO(shape.image.blob)).convert("RGBA")
    img = Image.alpha_composite(Image.new("RGBA", img.size, "white"), img).convert("RGB")
    d = ImageDraw.Draw(img)
    d.rectangle((205, 46, 248, 73), fill="white")        # the drawing's [M_i]
    d.rectangle((70, 167, 99, 191), fill="white")        # and its X_v
    arr = np.asarray(img)

    with plt.rc_context(STYLE):
        fig, ax = plt.subplots(figsize=(8.4, 5.6))
        ax.imshow(arr, extent=(0, arr.shape[1], arr.shape[0], 0))
        ax.set_xlim(-300, arr.shape[1] + 300)
        ax.set_ylim(arr.shape[0] + 10, -70)
        ax.axis("off")
        # feed: into the top of the vessel
        ax.annotate("feed: $F$ (m³/s)\nsubstrate $S_F$ (mol/m³)", xy=(70, 92), xytext=(-290, -20),
                    fontsize=15, color=BLUE,
                    arrowprops=dict(arrowstyle="-|>", color=BLUE, lw=2.6, mutation_scale=22))
        # outflow: a pipe through the right wall at the liquid level
        ax.plot([240, 268], [178, 178], color=INK, lw=2)
        ax.plot([240, 268], [194, 194], color=INK, lw=2)
        ax.annotate("", xy=(380, 186), xytext=(262, 186),
                    arrowprops=dict(arrowstyle="-|>", color=ORANGE, lw=2.6, mutation_scale=22))
        ax.text(392, 186, "outflow: $F$ (m³/s)\nwith $X$ and $S$", fontsize=15, color=ORANGE,
                va="center")
        ax.annotate("cells: biomass $X$ (kg/m³)", xy=(205, 215), xytext=(300, 80), fontsize=15,
                    arrowprops=dict(arrowstyle="->", color=INK, lw=1.4))
        ax.annotate("substrate $S$ (mol/m³)", xy=(135, 285), xytext=(300, 300), fontsize=15,
                    arrowprops=dict(arrowstyle="->", color=INK, lw=1.4))
        ax.annotate("volume $V$ (m³)\nconstant: in = out", xy=(34, 178), xytext=(-290, 290),
                    fontsize=15, arrowprops=dict(arrowstyle="->", color=INK, lw=1.4))
        save(fig, "a05-reactor.png")


def data():
    rows = {}
    with open(DATA, newline="") as f:
        for r in csv.DictReader(f):
            rows.setdefault(int(r["experiment"]), []).append((float(r["t"]), float(r["X"]), float(r["S"])))
    colors = {1: "#1f77b4", 2: "#2ca02c", 3: "#9467bd", 4: "#e377c2"}

    with plt.rc_context(STYLE):
        fig, axes = plt.subplots(1, 2, figsize=(13, 4.8))
        for ax, j, name, unit, ss in [(axes[0], 1, "biomass $X$", "kg/m³", X_SS),
                                      (axes[1], 2, "substrate $S$", "mol/m³", S_SS)]:
            ax.axvspan(0, 100, color="#1f77b4", alpha=0.07, lw=0)
            for exp, vals in rows.items():
                v = np.array(vals)
                if exp == 5:
                    ax.plot(v[:, 0], v[:, j], "-o", color=INK, lw=0.8, mfc="white", mec=INK,
                            ms=5, mew=1.3, label="experiment 5 (test)", zorder=3)
                    first = v[:, 0] <= 40
                    ax.plot(v[first, 0], v[first, j], "o", color=INK, ms=5, zorder=4,
                            label="experiment 5, first 40 s\n(Problems 1 and 2 train on it)")
                else:
                    ax.plot(v[:, 0], v[:, j], "-o", color=colors[exp], lw=0.9, ms=3.5,
                            label=f"experiment {exp}")
            ax.axhline(ss, color=MUTED, lw=1.2, ls="--", label="steady state", zorder=0)
            ax.axvline(40, color=INK, lw=0.8, ls=":", zorder=0)
            ax.set_xlabel("time $t$ (s)")
            ax.set_ylabel(f"{name} ({unit})")
            ax.set_xlim(-5, 305)
            lo, hi = ax.get_ylim()
            ax.set_ylim(lo, hi + 0.22 * (hi - lo))
            top = hi + 0.2 * (hi - lo)
            ax.text(50, top, "training data\n(experiments 1 to 4)", ha="center", va="top",
                    fontsize=12, color="#1f5c99")
            ax.text(200, top, "beyond the training data\n(experiment 5 only)", ha="center",
                    va="top", fontsize=12, color=INK)
        axes[1].legend(loc="center right", fontsize=11, ncol=1)
        fig.tight_layout()
        save(fig, "a05-data.png")


GREEN = "#2e7d32"
L13_FIGURES = REPO / "lectures" / "l13" / "figures"

# Lecture 13's ways to add the physics to a network, as on its slide "Scientific machine learning,
# four ways to add the physics": the two this homework uses, the PINN and the neural ODE.
CARDS = [
    ("PINN", "pinn-vs-nn.png", r"$L = L_{\mathrm{data}} + \lambda\, L_{\mathrm{physics}}$",
     "physics in the loss:\na penalty, small, not zero", "Problem 2"),
    ("Neural ODE", "card-neural-ode.png", r"$\dfrac{dx}{dt} = f\left(x,\ \mathrm{NN}(x;\theta)\right)$",
     "physics in the equations:\nthe network sits inside them", "Problem 4"),
]


def physics():
    from matplotlib.patches import FancyBboxPatch

    with plt.rc_context(STYLE):
        fig = plt.figure(figsize=(12, 5.6))
        w = 0.48                                   # width of one card, in figure units
        for k, (title, img, formula, where, problem) in enumerate(CARDS):
            x0 = 0.01 + k * (w + 0.02)
            fig.patches.append(FancyBboxPatch(
                (x0, 0.08), w, 0.9,
                boxstyle="round,pad=0.004,rounding_size=0.015",
                transform=fig.transFigure,
                facecolor="white",
                edgecolor=BLUE,
                linewidth=3.5,
                zorder=-10,
            ))
            ax = fig.add_axes([x0 + 0.02, 0.47, w - 0.04, 0.40])
            ax.set_zorder(5)
            ax.imshow(plt.imread(L13_FIGURES / img))
            ax.axis("off")
            fig.text(x0 + w / 2, 0.925, title, ha="center", va="center", fontsize=18, fontweight="bold")
            fig.text(x0 + w / 2, 0.38, formula, ha="center", va="center", fontsize=16)
            fig.text(x0 + w / 2, 0.255, where, ha="center", va="center", fontsize=14, color=GREEN)
            fig.text(x0 + w / 2, 0.14, problem, ha="center", va="center", fontsize=15,
                     fontweight="bold", color=BLUE,
                     bbox=dict(boxstyle="round,pad=0.35", facecolor="#e8f0f8", edgecolor=BLUE))
        fig.text(0.5, 0.025, "Problems 1 and 3 use no physics: the neural network of time and the "
                 "time-series model learn from data only.", ha="center", va="center", fontsize=13,
                 color=MUTED)
        save(fig, "a05-physics.png")


if __name__ == "__main__":
    reactor()
    data()
    physics()
